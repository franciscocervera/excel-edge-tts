#!/usr/bin/env python3

import argparse
import asyncio
import re
from collections import Counter
from pathlib import Path

import edge_tts
from openpyxl import load_workbook


DEFAULT_VOICE = "es-MX-DaliaNeural"
DEFAULT_SHEET = "Audios_TTS"
DEFAULT_RATE = "+0%"
DEFAULT_VOLUME = "+0%"
DEFAULT_PITCH = "+0Hz"
DEFAULT_CONCURRENCY = 3
DEFAULT_RETRIES = 3


def sanitize_filename(value: str) -> str:
    value = value.strip()
    value = re.sub(r'[<>:"/\\|?*\x00-\x1F]', "_", value)
    return value.rstrip(". ")


def read_audio_rows(excel_path: Path, sheet_name: str) -> list[tuple[str, str]]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)

    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"No existe la hoja '{sheet_name}'. "
                f"Hojas disponibles: {', '.join(workbook.sheetnames)}"
            )

        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)

        try:
            headers = next(rows)
        except StopIteration as exc:
            raise ValueError("El Excel está vacío.") from exc

        header_map = {
            str(value).strip(): index
            for index, value in enumerate(headers)
            if value is not None
        }

        required = {"Audio_ID", "Texto"}
        missing = required - set(header_map)
        if missing:
            raise ValueError(
                "Faltan columnas requeridas: " + ", ".join(sorted(missing))
            )

        audio_rows: list[tuple[str, str]] = []

        for row_number, row in enumerate(rows, start=2):
            audio_id = row[header_map["Audio_ID"]]
            text = row[header_map["Texto"]]

            if audio_id is None and text is None:
                continue

            if audio_id is None or not str(audio_id).strip():
                print(f"[OMITIDO] Fila {row_number}: Audio_ID vacío.")
                continue

            if text is None or not str(text).strip():
                print(f"[OMITIDO] {audio_id}: texto vacío.")
                continue

            audio_rows.append((str(audio_id).strip(), str(text).strip()))

        return audio_rows
    finally:
        workbook.close()


def validate_audio_ids(audio_rows: list[tuple[str, str]]) -> None:
    ids = [audio_id for audio_id, _ in audio_rows]
    duplicates = sorted(
        audio_id for audio_id, count in Counter(ids).items() if count > 1
    )
    if duplicates:
        raise ValueError("Audio_ID duplicados: " + ", ".join(duplicates))

    filenames: dict[str, str] = {}
    for audio_id in ids:
        filename = sanitize_filename(audio_id)
        if not filename:
            raise ValueError(f"Audio_ID inválido para nombre de archivo: {audio_id!r}")

        key = filename.casefold()
        previous = filenames.get(key)
        if previous and previous != audio_id:
            raise ValueError(
                "Audio_ID distintos generan el mismo nombre de archivo: "
                f"'{previous}' y '{audio_id}'."
            )
        filenames[key] = audio_id


async def generate_audio(
    audio_id: str,
    text: str,
    output_dir: Path,
    semaphore: asyncio.Semaphore,
    voice: str,
    rate: str,
    volume: str,
    pitch: str,
    overwrite: bool,
    retries: int,
) -> tuple[str, str]:
    filename = sanitize_filename(audio_id) + ".mp3"
    output_path = output_dir / filename

    if output_path.exists() and not overwrite:
        print(f"[EXISTE]  {filename}")
        return audio_id, "exists"

    async with semaphore:
        for attempt in range(1, retries + 1):
            try:
                communicate = edge_tts.Communicate(
                    text=text,
                    voice=voice,
                    rate=rate,
                    volume=volume,
                    pitch=pitch,
                )
                await communicate.save(str(output_path))

                if not output_path.exists() or output_path.stat().st_size == 0:
                    raise RuntimeError("El archivo generado está vacío.")

                print(f"[OK]      {filename}")
                return audio_id, "generated"

            except Exception as exc:
                output_path.unlink(missing_ok=True)

                if attempt == retries:
                    print(f"[ERROR]   {filename}: {exc}")
                    return audio_id, "error"

                print(f"[REINTENTO {attempt}/{retries}] {filename}: {exc}")
                await asyncio.sleep(min(2**attempt, 8))

    return audio_id, "error"


async def main_async(args: argparse.Namespace) -> int:
    excel_path = Path(args.excel).expanduser().resolve()
    output_dir = Path(args.output).expanduser().resolve()

    if not excel_path.exists():
        print(f"No existe el archivo: {excel_path}")
        return 1

    if not excel_path.is_file():
        print(f"La ruta del Excel no es un archivo: {excel_path}")
        return 1

    output_dir.mkdir(parents=True, exist_ok=True)

    try:
        audio_rows = read_audio_rows(excel_path, args.sheet)
        validate_audio_ids(audio_rows)
    except Exception as exc:
        print(f"Error al leer el catálogo: {exc}")
        return 1

    if not audio_rows:
        print("No se encontraron audios para generar.")
        return 0

    print(f"Excel:       {excel_path}")
    print(f"Hoja:        {args.sheet}")
    print(f"Voz:         {args.voice}")
    print(f"Audios:      {len(audio_rows)}")
    print(f"Destino:     {output_dir}")
    print(f"Velocidad:   {args.rate}")
    print(f"Volumen:     {args.volume}")
    print(f"Tono:        {args.pitch}")
    print()

    semaphore = asyncio.Semaphore(args.concurrency)
    tasks = [
        generate_audio(
            audio_id=audio_id,
            text=text,
            output_dir=output_dir,
            semaphore=semaphore,
            voice=args.voice,
            rate=args.rate,
            volume=args.volume,
            pitch=args.pitch,
            overwrite=args.overwrite,
            retries=args.retries,
        )
        for audio_id, text in audio_rows
    ]

    results = await asyncio.gather(*tasks)

    generated = sum(status == "generated" for _, status in results)
    existing = sum(status == "exists" for _, status in results)
    errors = sum(status == "error" for _, status in results)

    print()
    print("Resumen")
    print(f"  Generados: {generated}")
    print(f"  Existían:  {existing}")
    print(f"  Errores:   {errors}")

    return 1 if errors else 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Genera archivos MP3 en lote desde un catálogo Excel usando edge-tts."
        )
    )
    parser.add_argument(
        "excel",
        nargs="?",
        default="catalogo_audios.xlsx",
        help="Ruta del archivo Excel. Predeterminado: catalogo_audios.xlsx",
    )
    parser.add_argument(
        "--sheet",
        default=DEFAULT_SHEET,
        help=f"Hoja del Excel. Predeterminado: {DEFAULT_SHEET}",
    )
    parser.add_argument(
        "--output",
        default="audio",
        help="Carpeta de salida. Predeterminado: audio",
    )
    parser.add_argument(
        "--voice",
        default=DEFAULT_VOICE,
        help=f"Voz de Edge TTS. Predeterminado: {DEFAULT_VOICE}",
    )
    parser.add_argument(
        "--rate",
        default=DEFAULT_RATE,
        help='Velocidad, por ejemplo "+0%%" o "-5%%".',
    )
    parser.add_argument(
        "--volume",
        default=DEFAULT_VOLUME,
        help='Volumen, por ejemplo "+0%%" o "+10%%".',
    )
    parser.add_argument(
        "--pitch",
        default=DEFAULT_PITCH,
        help='Tono, por ejemplo "+0Hz" o "-5Hz".',
    )
    parser.add_argument(
        "--concurrency",
        type=int,
        default=DEFAULT_CONCURRENCY,
        help=f"Solicitudes simultáneas. Predeterminado: {DEFAULT_CONCURRENCY}",
    )
    parser.add_argument(
        "--retries",
        type=int,
        default=DEFAULT_RETRIES,
        help=f"Reintentos por audio. Predeterminado: {DEFAULT_RETRIES}",
    )
    parser.add_argument(
        "--overwrite",
        action="store_true",
        help="Regenera también los MP3 existentes.",
    )
    return parser


def main() -> None:
    parser = build_parser()
    args = parser.parse_args()

    if args.concurrency < 1:
        parser.error("--concurrency debe ser mayor o igual a 1.")

    if args.retries < 1:
        parser.error("--retries debe ser mayor o igual a 1.")

    raise SystemExit(asyncio.run(main_async(args)))


if __name__ == "__main__":
    main()
